"""Distribuisce i dati delle sezioni di censimento Istat 2021 (provincia 108) su una griglia esagonale -> hex_mb.json.
Richiede: shapely, pyshp, openpyxl (PYTHONPATH=C:/Temp/pl). File: istat_samples/sezioni/x/SHP/R03_21_WGS84.*, istat_samples/census_dl/R03_indicatori_2021_sezioni.xlsx
Ogni valore di sezione e' ripartito sugli esagoni in proporzione all'area di intersezione (ipotesi: densita' uniforme nella sezione)."""
import json, math, sys
import numpy as np
import openpyxl, shapefile, shapely
from pyproj import Transformer
from shapely.geometry import Polygon, shape
from shapely.strtree import STRtree

LON0, LAT0 = 9.27, 45.62
KX, KY = 111.32 * math.cos(math.radians(LAT0)), 110.57
R = 0.45  # raggio esagono in km
T = Transformer.from_crs(32632, 4326, always_xy=True)  # il file 'WGS84' di Istat e' in UTM 32N (metri)


def tr(a):
    lon, lat = T.transform(a[:, 0], a[:, 1])
    return np.column_stack(((lon - LON0) * KX, -(lat - LAT0) * KY))

# indicatori dal file xlsx: P101 occupati 15-64, ST1 stranieri, P90 laureati e terziario
wb = openpyxl.load_workbook("istat_samples/census_dl/R03_indicatori_2021_sezioni.xlsx", read_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(min_row=1, max_row=70000, values_only=True)
h = next(rows)
ix = {k: h.index(k) for k in ("SEZ21_ID", "PROCOM", "P1", "P101", "ST1", "P90", "A8")}
ind = {}
for n, r in enumerate(rows):
    if n % 10000 == 0: print("xlsx riga", n, file=sys.stderr, flush=True)
    if r[ix["PROCOM"]] is None: continue
    if str(r[ix["PROCOM"]]).startswith("108"):
        ind[int(r[ix["SEZ21_ID"]])] = {k: (r[ix[k]] or 0) for k in ("P1", "P101", "ST1", "P90", "A8")}
print("sezioni con indicatori:", len(ind), file=sys.stderr)

sf = shapefile.Reader("istat_samples/sezioni/x/SHP/R03_21_WGS84")
secs = []
for sr in sf.iterShapeRecords():
    d = sr.record.as_dict()
    if not str(d["PRO_COM"]).startswith("108"): continue
    g = shape(sr.shape.__geo_interface__)
    if g.is_empty: continue
    g = type(g)(g) if False else g
    g = shapely.transform(g, tr)
    v = ind.get(int(d["SEZ21_ID"]), {})
    secs.append((str(d["PRO_COM"]), g.buffer(0), {"pop": d["POP21"], "emp": v.get("P101", 0), "for": v.get("ST1", 0), "deg": v.get("P90", 0), "dw": d["ABI21"]}))
print("sezioni geometrie:", len(secs), file=sys.stderr)

minx = min(g.bounds[0] for _, g, _ in secs); miny = min(g.bounds[1] for _, g, _ in secs)
maxx = max(g.bounds[2] for _, g, _ in secs); maxy = max(g.bounds[3] for _, g, _ in secs)
w, hh = math.sqrt(3) * R, 1.5 * R
cells = []
row = 0
y = miny - R
while y < maxy + R:
    x = minx - R + (w / 2 if row % 2 else 0)
    while x < maxx + R:
        cells.append(Polygon([(x + R * math.cos(math.radians(30 + 60 * k)), y + R * math.sin(math.radians(30 + 60 * k))) for k in range(6)]))
        x += w
    y += hh; row += 1
centers = [(c.centroid.x, c.centroid.y) for c in cells]
tree = STRtree(cells)
keys = ("pop", "emp", "for", "deg", "dw")
acc = [dict.fromkeys(keys, 0.0) for _ in cells]
by, monza = {}, []
for pc, g, v in secs:
    if g.area <= 0: continue
    b = by.setdefault(pc, dict.fromkeys(keys, 0))
    for k in keys: b[k] += v[k]
    if pc == '108033': monza.append([round(g.centroid.x, 2), round(g.centroid.y, 2), int(v['pop'])])
    for i in tree.query(g):
        a = cells[i].intersection(g).area
        if a <= 0: continue
        f = a / g.area
        for k in keys: acc[i][k] += v[k] * f

out = {"r": R, "cells": []}
for (x, y), a in zip(centers, acc):
    if a["pop"] >= 1:
        out["cells"].append([round(x, 2), round(y, 2)] + [int(round(a[k])) for k in keys])
out["keys"] = list(keys)
out["by"] = {k: {kk: int(vv) for kk, vv in v.items()} for k, v in by.items()}
out["monza_sec"] = monza
out["totals"] = {k: int(round(sum(c[2 + i] for c in out["cells"]))) for i, k in enumerate(keys)}
out["note"] = "Istat Censimento 2021, sezioni di censimento ripartite su esagoni per area di intersezione"
json.dump(out, open("hex_mb.json", "w"), separators=(",", ":"))
print("esagoni:", len(out["cells"]), "totali:", out["totals"], file=sys.stderr)
